"""
tests/test_10-export-expenses-report.py

Pytest test suite for Step 10: Export Expenses Report (PDF + XLSX).

All assertions are driven from the feature spec at
.claude/specs/10-export-expenses-report.md — never from implementation details.

Fixture strategy mirrors tests/test_delete_expense.py:
- `app`           : Flask app wired to a temp-file SQLite DB via monkeypatch.
- `client`        : Plain (unauthenticated) test client.
- `db_conn`       : Direct SQLite connection to the same temp DB.
- `test_user_id`  : One owner user, inserted fresh per test.
- `other_user_id` : A second user for cross-user isolation tests.
- `auth_client`   : Client with `test_user_id` in the session.
"""

import io
import sqlite3

import openpyxl
import pytest

from app import app as flask_app
from database.db import init_db
from database.queries import get_expenses_for_export, insert_expense
from werkzeug.security import generate_password_hash


# ------------------------------------------------------------------ #
# Core fixtures                                                        #
# ------------------------------------------------------------------ #


@pytest.fixture
def app(tmp_path, monkeypatch):
    db_file = str(tmp_path / "test_export.db")
    import database.db as db_module

    monkeypatch.setattr(db_module, "DB_PATH", db_file)

    flask_app.config.update(
        {
            "TESTING": True,
            "SECRET_KEY": "test-secret",
            "WTF_CSRF_ENABLED": False,
        }
    )

    with flask_app.app_context():
        init_db()
        yield flask_app


@pytest.fixture
def client(app):
    return app.test_client()


@pytest.fixture
def db_conn(app):
    import database.db as db_module

    conn = sqlite3.connect(db_module.DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    yield conn
    conn.commit()
    conn.close()


def _create_user(db_conn, name, email):
    cur = db_conn.execute(
        "INSERT INTO users (name, email, password_hash) VALUES (?, ?, ?)",
        (name, email, generate_password_hash("pw", method="pbkdf2:sha256")),
    )
    db_conn.commit()
    return cur.lastrowid


@pytest.fixture
def test_user_id(db_conn):
    return _create_user(db_conn, "Owner User", "owner@example.com")


@pytest.fixture
def other_user_id(db_conn):
    return _create_user(db_conn, "Other User", "other@example.com")


@pytest.fixture
def auth_client(client, test_user_id):
    with client.session_transaction() as sess:
        sess["user_id"] = test_user_id
        sess["user_name"] = "Owner User"
    return client


# ------------------------------------------------------------------ #
# Shared helpers                                                       #
# ------------------------------------------------------------------ #

XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _seed_owner_expenses(app, user_id):
    """Insert three expenses for `user_id` spanning different dates."""
    with app.app_context():
        insert_expense(user_id, 500.00, "Food", "2026-04-10", "Lunch A")
        insert_expense(user_id, 1200.00, "Bills", "2026-04-15", "Electric bill")
        insert_expense(user_id, 300.00, "Transport", "2026-05-01", "Cab ride")


def _seed_other_expenses(app, user_id):
    """Insert two expenses for a second user."""
    with app.app_context():
        insert_expense(
            user_id, 9999.00, "Shopping", "2026-05-10", "Other user purchase"
        )
        insert_expense(user_id, 7777.00, "Health", "2026-05-12", "Other user pharmacy")


def _workbook_from_response(response):
    """Parse openpyxl Workbook from a Flask test response."""
    return openpyxl.load_workbook(io.BytesIO(response.data))


# ------------------------------------------------------------------ #
# 1. Unit tests — get_expenses_for_export                             #
# ------------------------------------------------------------------ #


class TestGetExpensesForExportHelper:

    def test_returns_all_expenses_for_user_no_date_range(self, app, test_user_id):
        """User with expenses, no date range → all rows returned."""
        _seed_owner_expenses(app, test_user_id)
        with app.app_context():
            results = get_expenses_for_export(test_user_id)
        assert len(results) == 3, "Expected all 3 expenses with no date range"

    def test_returns_empty_list_for_user_with_no_expenses(self, app, test_user_id):
        """User with no expenses → empty list, not an error."""
        with app.app_context():
            results = get_expenses_for_export(test_user_id)
        assert results == [], "Expected [] for a user with no expenses"

    def test_returns_only_expenses_inside_valid_date_range(self, app, test_user_id):
        """Valid date range → only expenses whose date falls inside the range."""
        _seed_owner_expenses(app, test_user_id)
        with app.app_context():
            results = get_expenses_for_export(
                test_user_id, date_from="2026-04-01", date_to="2026-04-30"
            )
        # Only the two April expenses should appear
        assert len(results) == 2, "Expected 2 expenses in the April range"
        for row in results:
            assert row["date"] >= "2026-04-01"
            assert row["date"] <= "2026-04-30"

    def test_returns_empty_list_for_range_that_excludes_everything(
        self, app, test_user_id
    ):
        """Date range with no matching expenses → empty list."""
        _seed_owner_expenses(app, test_user_id)
        with app.app_context():
            results = get_expenses_for_export(
                test_user_id, date_from="2020-01-01", date_to="2020-01-31"
            )
        assert results == [], "Expected [] when date range excludes all expenses"

    def test_returns_only_that_users_data_when_called_with_other_user_id(
        self, app, test_user_id, other_user_id
    ):
        """Calling with another user's id must return only that user's data."""
        _seed_owner_expenses(app, test_user_id)
        _seed_other_expenses(app, other_user_id)
        with app.app_context():
            results = get_expenses_for_export(other_user_id)
        amounts = {r["amount"] for r in results}
        descriptions = {r["description"] for r in results}
        # Must not contain the owner user's data
        assert 500.00 not in amounts, "Must not contain owner user's expenses"
        assert "Lunch A" not in descriptions
        # Must contain only the other user's data
        assert 9999.00 in amounts
        assert "Other user purchase" in descriptions

    def test_only_date_from_provided_falls_back_to_all_time(self, app, test_user_id):
        """Only date_from supplied (no date_to) → both-or-neither contract → all expenses."""
        _seed_owner_expenses(app, test_user_id)
        with app.app_context():
            # Passing only date_from — date_to stays None
            results = get_expenses_for_export(
                test_user_id, date_from="2026-04-01", date_to=None
            )
        # Both-or-neither: when one side is missing, no date filter is applied
        assert (
            len(results) == 3
        ), "With only date_from provided the helper must return all expenses (both-or-neither)"

    def test_results_are_ordered_newest_first(self, app, test_user_id):
        """Results must be ordered date DESC (newest first)."""
        _seed_owner_expenses(app, test_user_id)
        with app.app_context():
            results = get_expenses_for_export(test_user_id)
        dates = [r["date"] for r in results]
        assert dates == sorted(
            dates, reverse=True
        ), "Results must be ordered newest first (date DESC)"


# ------------------------------------------------------------------ #
# 2. PDF route tests                                                   #
# ------------------------------------------------------------------ #


class TestPdfRouteAuthGuard:

    def test_unauthenticated_returns_302(self, client):
        response = client.get("/expenses/export/pdf")
        assert (
            response.status_code == 302
        ), "Unauthenticated request must redirect (302)"

    def test_unauthenticated_redirects_to_login(self, client):
        response = client.get("/expenses/export/pdf")
        assert "/login" in response.headers.get(
            "Location", ""
        ), "Unauthenticated request must redirect to /login"


class TestPdfRouteHappyPath:

    def test_authenticated_returns_200(self, app, auth_client, test_user_id):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf")
        assert response.status_code == 200, "Authenticated export must return 200"

    def test_response_content_type_is_pdf(self, app, auth_client, test_user_id):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf")
        assert response.content_type.startswith(
            "application/pdf"
        ), f"Expected application/pdf, got {response.content_type}"

    def test_response_disposition_is_attachment(self, app, auth_client, test_user_id):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf")
        disposition = response.headers.get("Content-Disposition", "")
        assert (
            "attachment" in disposition
        ), "Content-Disposition must contain 'attachment'"

    def test_response_disposition_contains_pdf_extension(
        self, app, auth_client, test_user_id
    ):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf")
        disposition = response.headers.get("Content-Disposition", "")
        assert (
            ".pdf" in disposition
        ), "Content-Disposition filename must include .pdf extension"

    def test_response_body_starts_with_pdf_magic_bytes(
        self, app, auth_client, test_user_id
    ):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf")
        assert (
            response.data[:5] == b"%PDF-"
        ), "PDF response body must start with the %PDF- magic bytes"

    def test_all_time_filename_contains_all_time_label(
        self, app, auth_client, test_user_id
    ):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf")
        disposition = response.headers.get("Content-Disposition", "")
        assert (
            "all-time" in disposition
        ), "Filename for an all-time export must contain 'all-time'"

    def test_all_time_filename_contains_slugified_user_name(
        self, app, auth_client, test_user_id
    ):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf")
        disposition = response.headers.get("Content-Disposition", "")
        # "Owner User" slugifies to "owner-user"
        assert (
            "owner-user" in disposition
        ), "Filename must include a slugified version of the user's name"


class TestPdfRouteDateFilter:

    def test_filtered_export_returns_200(self, app, auth_client, test_user_id):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get(
            "/expenses/export/pdf?date_from=2026-04-01&date_to=2026-04-30"
        )
        assert response.status_code == 200

    def test_filtered_filename_includes_date_range(
        self, app, auth_client, test_user_id
    ):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get(
            "/expenses/export/pdf?date_from=2026-04-01&date_to=2026-04-30"
        )
        disposition = response.headers.get("Content-Disposition", "")
        assert (
            "2026-04-01" in disposition
        ), "Filtered export filename must include date_from"
        assert (
            "2026-04-30" in disposition
        ), "Filtered export filename must include date_to"

    def test_filtered_filename_does_not_contain_all_time(
        self, app, auth_client, test_user_id
    ):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get(
            "/expenses/export/pdf?date_from=2026-04-01&date_to=2026-04-30"
        )
        disposition = response.headers.get("Content-Disposition", "")
        assert (
            "all-time" not in disposition
        ), "Filename with a specific range must not say 'all-time'"


class TestPdfRouteZeroExpenses:

    def test_zero_expenses_returns_200(self, auth_client):
        """No expenses in the DB → must still return 200, not 404 or 500."""
        response = auth_client.get("/expenses/export/pdf")
        assert (
            response.status_code == 200
        ), "Zero-expense export must return 200, not 404 or 500"

    def test_zero_expenses_returns_valid_pdf(self, auth_client):
        """No expenses → well-formed PDF (starts with %PDF-)."""
        response = auth_client.get("/expenses/export/pdf")
        assert (
            response.data[:5] == b"%PDF-"
        ), "Zero-expense export must still return a valid PDF"

    def test_zero_expenses_content_type_is_pdf(self, auth_client):
        response = auth_client.get("/expenses/export/pdf")
        assert response.content_type.startswith("application/pdf")


class TestPdfRouteInvalidDateFallback:

    def test_garbage_date_from_does_not_500(self, app, auth_client, test_user_id):
        """Invalid date_from must fall back to all-time — never a 500."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf?date_from=garbage")
        assert (
            response.status_code == 200
        ), "Garbage date_from must silently fall back to all-time, not 500"

    def test_garbage_date_from_filename_says_all_time(
        self, app, auth_client, test_user_id
    ):
        """When date_from is invalid the filename must fall back to all-time."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf?date_from=garbage")
        disposition = response.headers.get("Content-Disposition", "")
        assert (
            "all-time" in disposition
        ), "Invalid date falls back to all-time; filename must reflect this"

    def test_only_date_from_valid_no_date_to_falls_back_to_all_time(
        self, app, auth_client, test_user_id
    ):
        """Only one side of the pair → both-or-neither → all-time fallback."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/pdf?date_from=2026-04-01")
        assert response.status_code == 200
        disposition = response.headers.get("Content-Disposition", "")
        assert "all-time" in disposition


# ------------------------------------------------------------------ #
# 3. XLSX route tests                                                  #
# ------------------------------------------------------------------ #


class TestXlsxRouteAuthGuard:

    def test_unauthenticated_returns_302(self, client):
        response = client.get("/expenses/export/xlsx")
        assert (
            response.status_code == 302
        ), "Unauthenticated request must redirect (302)"

    def test_unauthenticated_redirects_to_login(self, client):
        response = client.get("/expenses/export/xlsx")
        assert "/login" in response.headers.get(
            "Location", ""
        ), "Unauthenticated request must redirect to /login"


class TestXlsxRouteHappyPath:

    def test_authenticated_returns_200(self, app, auth_client, test_user_id):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        assert response.status_code == 200

    def test_response_content_type_is_xlsx(self, app, auth_client, test_user_id):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        assert (
            response.content_type == XLSX_MIMETYPE
        ), f"Expected exact xlsx mimetype, got {response.content_type}"

    def test_response_disposition_is_attachment(self, app, auth_client, test_user_id):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        disposition = response.headers.get("Content-Disposition", "")
        assert "attachment" in disposition

    def test_response_disposition_contains_xlsx_extension(
        self, app, auth_client, test_user_id
    ):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        disposition = response.headers.get("Content-Disposition", "")
        assert (
            ".xlsx" in disposition
        ), "Content-Disposition filename must include .xlsx extension"

    def test_workbook_is_openpyxl_parseable(self, app, auth_client, test_user_id):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        # Should not raise
        wb = _workbook_from_response(response)
        assert wb is not None

    def test_workbook_has_sheet_named_expenses(self, app, auth_client, test_user_id):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        wb = _workbook_from_response(response)
        assert (
            "Expenses" in wb.sheetnames
        ), f"Expected a sheet named 'Expenses', found: {wb.sheetnames}"

    def test_header_row_matches_spec(self, app, auth_client, test_user_id):
        """Row 1 must be exactly ['Date', 'Description', 'Category', 'Amount (INR)']."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        wb = _workbook_from_response(response)
        ws = wb["Expenses"]
        header = [ws.cell(row=1, column=c).value for c in range(1, 5)]
        assert header == [
            "Date",
            "Description",
            "Category",
            "Amount (INR)",
        ], f"Header row mismatch: {header}"

    def test_amount_cells_are_numeric(self, app, auth_client, test_user_id):
        """Amount (INR) column values in data rows must be numeric, not strings."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        wb = _workbook_from_response(response)
        ws = wb["Expenses"]
        max_row = ws.max_row
        # Data rows start at row 2; skip the last row (Total) by checking data rows only
        # We have 3 data rows + 1 header + 1 total = 5 rows total
        for row_idx in range(2, max_row):  # exclude the Total row
            cell = ws.cell(row=row_idx, column=4)
            # openpyxl reads formula cells as None or a string; a plain numeric value is int/float
            # The spec requires numeric cells — they must not be plain strings
            assert not isinstance(
                cell.value, str
            ), f"Amount cell at row {row_idx} must be numeric, not a string: {cell.value!r}"

    def test_total_row_label_is_present(self, app, auth_client, test_user_id):
        """The last row must have 'Total' in the first column."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        wb = _workbook_from_response(response)
        ws = wb["Expenses"]
        last_row = ws.max_row
        total_label = ws.cell(row=last_row, column=1).value
        assert (
            total_label == "Total"
        ), f"Expected 'Total' label in last row column A, got: {total_label!r}"

    def test_total_row_formula_references_sum(self, app, auth_client, test_user_id):
        """The last row's Amount cell must contain a SUM formula."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        # data_only=False to retrieve formulas rather than cached values
        wb = openpyxl.load_workbook(io.BytesIO(response.data), data_only=False)
        ws = wb["Expenses"]
        last_row = ws.max_row
        formula_cell = ws.cell(row=last_row, column=4)
        formula_value = formula_cell.value
        assert isinstance(formula_value, str) and formula_value.upper().startswith(
            "=SUM("
        ), f"Expected a SUM formula in the Total row, got: {formula_value!r}"

    def test_all_time_filename_contains_all_time_label(
        self, app, auth_client, test_user_id
    ):
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        disposition = response.headers.get("Content-Disposition", "")
        assert "all-time" in disposition


class TestXlsxRouteZeroExpenses:

    def test_zero_expenses_returns_200(self, auth_client):
        """No expenses → still 200, never 404 or 500."""
        response = auth_client.get("/expenses/export/xlsx")
        assert response.status_code == 200, "Zero-expense xlsx export must return 200"

    def test_zero_expenses_workbook_is_well_formed(self, auth_client):
        """No expenses → workbook still parseable by openpyxl."""
        response = auth_client.get("/expenses/export/xlsx")
        wb = _workbook_from_response(response)
        assert "Expenses" in wb.sheetnames

    def test_zero_expenses_header_row_present(self, auth_client):
        """No expenses → header row must still be present."""
        response = auth_client.get("/expenses/export/xlsx")
        wb = _workbook_from_response(response)
        ws = wb["Expenses"]
        header = [ws.cell(row=1, column=c).value for c in range(1, 5)]
        assert header == [
            "Date",
            "Description",
            "Category",
            "Amount (INR)",
        ], "Header row must be present even with zero expenses"

    def test_zero_expenses_no_data_rows_below_header(self, auth_client):
        """No expenses → only the header row; no data rows below it."""
        response = auth_client.get("/expenses/export/xlsx")
        wb = _workbook_from_response(response)
        ws = wb["Expenses"]
        # When there are zero expenses the spec says no Total row is needed,
        # so there should be no second row with actual expense data.
        row2_col1 = ws.cell(row=2, column=1).value
        # Either no row exists (None) or it's not a date string — not an expense row
        if row2_col1 is not None:
            # Could be None or "Total" but must not look like a data row date
            assert not str(row2_col1).startswith(
                "2026"
            ), "No data rows should be present for a zero-expense workbook"


# ------------------------------------------------------------------ #
# 4. Cross-user isolation tests                                        #
# ------------------------------------------------------------------ #


class TestCrossUserIsolation:
    """Logged-in user A must never see user B's data in any export."""

    def test_pdf_export_excludes_other_users_amounts(
        self, app, auth_client, test_user_id, other_user_id
    ):
        """User A's PDF must not contain user B's expense amounts."""
        _seed_owner_expenses(app, test_user_id)
        _seed_other_expenses(app, other_user_id)
        response = auth_client.get("/expenses/export/pdf")
        body = response.data
        # User B's amounts are 9999.00 and 7777.00 — these must not appear
        assert (
            b"9,999.00" not in body and b"9999" not in body
        ), "User A's PDF must not contain user B's 9999 amount"
        assert (
            b"7,777.00" not in body and b"7777" not in body
        ), "User A's PDF must not contain user B's 7777 amount"

    def test_pdf_export_excludes_other_users_descriptions(
        self, app, auth_client, test_user_id, other_user_id
    ):
        """User A's PDF must not contain user B's expense descriptions."""
        _seed_owner_expenses(app, test_user_id)
        _seed_other_expenses(app, other_user_id)
        response = auth_client.get("/expenses/export/pdf")
        body = response.data
        assert (
            b"Other user purchase" not in body
        ), "User A's PDF must not contain user B's description"
        assert b"Other user pharmacy" not in body

    def test_xlsx_export_excludes_other_users_data(
        self, app, auth_client, test_user_id, other_user_id
    ):
        """User A's XLSX must contain only user A's rows."""
        _seed_owner_expenses(app, test_user_id)
        _seed_other_expenses(app, other_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        wb = _workbook_from_response(response)
        ws = wb["Expenses"]
        # Collect all cell values from data rows (row 2 onwards)
        all_values = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            all_values.extend([v for v in row if v is not None])
        # User B's amounts must not appear
        assert (
            9999.0 not in all_values and 9999 not in all_values
        ), "User A's XLSX must not contain user B's 9999 amount"
        assert (
            7777.0 not in all_values and 7777 not in all_values
        ), "User A's XLSX must not contain user B's 7777 amount"
        # User B's descriptions must not appear
        assert (
            "Other user purchase" not in all_values
        ), "User A's XLSX must not contain user B's description"
        assert "Other user pharmacy" not in all_values

    def test_xlsx_export_contains_only_owner_user_data(
        self, app, auth_client, test_user_id, other_user_id
    ):
        """User A's XLSX data rows must include only user A's own descriptions."""
        _seed_owner_expenses(app, test_user_id)
        _seed_other_expenses(app, other_user_id)
        response = auth_client.get("/expenses/export/xlsx")
        wb = _workbook_from_response(response)
        ws = wb["Expenses"]
        descriptions = [
            ws.cell(row=r, column=2).value
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=2).value not in (None, "Total")
        ]
        assert (
            "Lunch A" in descriptions
            or "Electric bill" in descriptions
            or "Cab ride" in descriptions
        ), "User A's own expenses must appear in the XLSX"


# ------------------------------------------------------------------ #
# 5. Profile page template tests                                       #
# ------------------------------------------------------------------ #


class TestProfilePageExportLinks:

    def test_profile_contains_pdf_export_link(self, app, auth_client, test_user_id):
        """Profile page must render a link to the PDF export route."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/profile")
        assert response.status_code == 200
        body = response.data.decode("utf-8")
        assert (
            'href="/expenses/export/pdf"' in body or "/expenses/export/pdf" in body
        ), "Profile page must contain a link to /expenses/export/pdf"

    def test_profile_contains_xlsx_export_link(self, app, auth_client, test_user_id):
        """Profile page must render a link to the XLSX export route."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/profile")
        assert response.status_code == 200
        body = response.data.decode("utf-8")
        assert (
            "/expenses/export/xlsx" in body
        ), "Profile page must contain a link to /expenses/export/xlsx"

    def test_profile_export_links_carry_date_filter_params_html_escaped(
        self, app, auth_client, test_user_id
    ):
        """When profile is loaded with date_from/date_to, both export links must
        carry those params. Jinja escapes & to &amp; in attribute values."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/profile?date_from=2026-04-01&date_to=2026-04-30")
        assert response.status_code == 200
        body = response.data.decode("utf-8")
        # Jinja2 HTML-escapes & to &amp; inside attribute values
        assert (
            "date_from=2026-04-01" in body
        ), "Export links must carry date_from parameter"
        assert "date_to=2026-04-30" in body, "Export links must carry date_to parameter"
        # The ampersand between params must appear HTML-escaped in the attribute
        assert (
            "date_from=2026-04-01&amp;date_to=2026-04-30" in body
        ), "Jinja must HTML-escape & to &amp; in the href attribute value"

    def test_profile_export_links_pdf_carries_date_filter(
        self, app, auth_client, test_user_id
    ):
        """The PDF link specifically must carry the active date filter."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/profile?date_from=2026-04-01&date_to=2026-04-30")
        body = response.data.decode("utf-8")
        # The PDF href with HTML-escaped params must appear
        assert "/expenses/export/pdf" in body
        # Verify the full href with params exists somewhere near the export/pdf route
        assert "export/pdf" in body and "date_from=2026-04-01" in body

    def test_profile_export_links_xlsx_carries_date_filter(
        self, app, auth_client, test_user_id
    ):
        """The XLSX link specifically must carry the active date filter."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/profile?date_from=2026-04-01&date_to=2026-04-30")
        body = response.data.decode("utf-8")
        assert "/expenses/export/xlsx" in body
        assert "export/xlsx" in body and "date_to=2026-04-30" in body

    def test_profile_all_time_view_has_no_empty_date_params(
        self, app, auth_client, test_user_id
    ):
        """When no date filter is active, export links must not include
        empty date_from=&date_to= strings in their hrefs."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/profile")
        body = response.data.decode("utf-8")
        assert (
            "date_from=&" not in body
        ), "All-time export links must not include empty date_from= params"
        assert (
            "date_from=&amp;" not in body
        ), "All-time export links must not include empty date_from=&amp; params"

    def test_profile_export_links_are_anchor_tags(self, app, auth_client, test_user_id):
        """Export links must be rendered as <a> tags, not buttons with JS."""
        _seed_owner_expenses(app, test_user_id)
        response = auth_client.get("/profile")
        body = response.data.decode("utf-8")
        # An <a ... href="...export/pdf..."> must exist
        assert "<a" in body
        assert "/expenses/export/pdf" in body
        assert "/expenses/export/xlsx" in body
