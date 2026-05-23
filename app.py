import html
import io
import math
import os
import re
import secrets
import sqlite3
from datetime import date, datetime, timedelta
from functools import wraps

from flask import (
    Flask,
    abort,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    session,
    send_file,
)
from werkzeug.security import check_password_hash

from database.db import init_db, seed_db
from database.queries import (
    CATEGORIES,
    create_user,
    get_user_by_email,
    get_user_by_id,
    get_summary_stats,
    get_recent_transactions,
    get_category_breakdown,
    insert_expense,
    get_expense_by_id,
    update_expense,
    delete_expense as delete_expense_query,
    get_expenses_for_export,
)

MAX_AMOUNT = 10_000_000  # ₹1 crore; reject larger inputs server-side
MIN_PASSWORD_LENGTH = 8

app = Flask(__name__)

_secret_key = os.environ.get("SECRET_KEY")
if not _secret_key:
    if os.environ.get("FLASK_ENV") == "production":
        raise RuntimeError("SECRET_KEY environment variable is required in production.")
    _secret_key = secrets.token_hex(32)
app.secret_key = _secret_key

app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    # Only force HTTPS-only cookies in production; dev runs on plain http://.
    SESSION_COOKIE_SECURE=os.environ.get("FLASK_ENV") == "production",
)


# ------------------------------------------------------------------ #
# CSRF protection                                                     #
# ------------------------------------------------------------------ #
# Hand-rolled to avoid an extra dependency. A random per-session token
# is generated lazily, injected into every Jinja render via the context
# processor below, and required on every unsafe HTTP method.


def _generate_csrf_token():
    if "_csrf_token" not in session:
        session["_csrf_token"] = secrets.token_hex(32)
    return session["_csrf_token"]


@app.before_request
def _csrf_protect():
    if app.testing:  # tests post forms without a token
        return
    if request.method not in ("POST", "PUT", "PATCH", "DELETE"):
        return
    expected = session.get("_csrf_token", "")
    submitted = request.form.get("csrf_token", "")
    if not expected or not secrets.compare_digest(expected, submitted):
        abort(400)


@app.context_processor
def _inject_csrf_token():
    return {"csrf_token": _generate_csrf_token}


# ------------------------------------------------------------------ #
# Auth guard                                                          #
# ------------------------------------------------------------------ #


def login_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapper


with app.app_context():
    init_db()
    seed_db()


# ------------------------------------------------------------------ #
# Date filter helpers (Step 6)                                        #
# ------------------------------------------------------------------ #


def _validate_iso_date(value):
    """Return value if it is a well-formed YYYY-MM-DD string, else None."""
    if not value:
        return None
    try:
        datetime.strptime(value, "%Y-%m-%d")
        return value
    except ValueError:
        return None


def _resolve_presets(today):
    """Build the bounded preset → (date_from, date_to) mapping for `today`.

    All Time has no params and is therefore not included here.
    """
    first_of_month = today.replace(day=1)
    return {
        "this_month": (first_of_month.isoformat(), today.isoformat()),
        "last_3_months": ((today - timedelta(days=90)).isoformat(), today.isoformat()),
        "last_6_months": ((today - timedelta(days=180)).isoformat(), today.isoformat()),
    }


def _validate_expense_form(amount_raw, category, date_raw):
    """Validate add/edit expense form fields. Return (amount_float, error_msg).

    `error_msg` is None on success; on failure `amount_float` is None.
    """
    try:
        amount = float(amount_raw)
    except ValueError:
        return None, "Amount must be a number greater than 0."
    if not math.isfinite(amount) or amount <= 0:
        return None, "Amount must be a number greater than 0."
    if amount > MAX_AMOUNT:
        return None, "Amount cannot exceed ₹1,00,00,000."
    if category not in CATEGORIES:
        return None, "Please choose a valid category."
    if _validate_iso_date(date_raw) is None:
        return None, "Please enter a valid date (YYYY-MM-DD)."
    return amount, None


# ------------------------------------------------------------------ #
# Export helpers (Step 10)                                            #
# ------------------------------------------------------------------ #

FONTS_DIR = os.path.join(os.path.dirname(__file__), "static", "fonts")
PDF_FONT_REGULAR = "DejaVuSans"
PDF_FONT_BOLD = "DejaVuSans-Bold"

XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Export design tokens. ReportLab/openpyxl can't read CSS variables, so we
# mirror the matching CSS tokens (--border, --paper, --ink-muted) here.
PDF_COLOR_HEADER_BG = "#e4e1da"
PDF_COLOR_ROW_ALT = "#f7f6f3"
PDF_COLOR_MUTED = "#6b6b6b"
XLSX_COLOR_HEADER_BG = "E4E1DA"  # openpyxl wants hex without the leading #

# Excel treats a cell that starts with one of these as a formula. Prefixing
# user-supplied text that starts with these characters with a leading
# apostrophe forces Excel/LibreOffice/Sheets to render it as a literal string
# instead of evaluating it (defeats =HYPERLINK("…") and similar payloads).
_XLSX_FORMULA_TRIGGERS = ("=", "+", "-", "@")


def _safe_xlsx_string(value):
    """Defang any string that Excel would interpret as a formula."""
    s = value or ""
    if s.startswith(_XLSX_FORMULA_TRIGGERS):
        return "'" + s
    return s


def _register_pdf_fonts():
    """Idempotently register the DejaVu fonts so the rupee glyph renders."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if PDF_FONT_REGULAR not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(
            TTFont(PDF_FONT_REGULAR, os.path.join(FONTS_DIR, "DejaVuSans.ttf"))
        )
        pdfmetrics.registerFont(
            TTFont(PDF_FONT_BOLD, os.path.join(FONTS_DIR, "DejaVuSans-Bold.ttf"))
        )


def _slugify_name(value):
    """ASCII-only, lowercase, hyphen-joined; safe to use in a filename."""
    ascii_only = (value or "").encode("ascii", "ignore").decode("ascii").lower()
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_only).strip("-")
    return slug or "user"


def _resolve_date_range_args(req):
    """Return (date_from, date_to, error) parsed from `date_from`/`date_to` args.

    Both-or-neither semantics: if exactly one side is supplied or one side
    fails to parse, the filter is dropped entirely. When both parse but the
    start is after the end we surface an error string so callers can flash
    it; exports just discard the range silently.
    """
    date_from = _validate_iso_date(req.args.get("date_from", "").strip())
    date_to = _validate_iso_date(req.args.get("date_to", "").strip())
    if (date_from is None) != (date_to is None):
        return None, None, None
    if date_from and date_to and date_from > date_to:
        return None, None, "Start date must be before end date."
    return date_from, date_to, None


# A paired concept: the same active range gets two presentations — a hyphenated
# slug for the download filename, and a more readable phrase for the PDF body.
def _filename_range_label(date_from, date_to):
    if date_from and date_to:
        return f"{date_from}-to-{date_to}"
    return "all-time"


def _human_range_label(date_from, date_to):
    if date_from and date_to:
        return f"{date_from} to {date_to}"
    return "All time"


def _expense_form_fields():
    """Trimmed dict of the four add/edit-expense form fields."""
    return {
        "amount": request.form.get("amount", "").strip(),
        "category": request.form.get("category", "").strip(),
        "date": request.form.get("date", "").strip(),
        "description": request.form.get("description", "").strip(),
    }


def _build_pdf(user_name, range_label, expenses):
    """Render the expense list as a PDF and return a seek-at-0 BytesIO."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )

    _register_pdf_fonts()

    # ReportLab's Paragraph parses XML; escape any user-controlled freetext
    # before interpolating it into a Paragraph string so a name like
    # `<img src="…"/>` cannot smuggle markup (or trigger an SSRF) at render time.
    safe_user_name = html.escape(user_name)

    total = sum(e["amount"] for e in expenses)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Title",
        parent=styles["Title"],
        fontName=PDF_FONT_BOLD,
        fontSize=18,
        leading=22,
        spaceAfter=12,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName=PDF_FONT_REGULAR,
        fontSize=11,
        leading=14,
    )
    empty_style = ParagraphStyle(
        "Empty",
        parent=body_style,
        fontName=PDF_FONT_REGULAR,
        textColor=colors.HexColor(PDF_COLOR_MUTED),
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title="Spendly Expense Report",
    )

    story = [
        Paragraph("Spendly Expense Report", title_style),
        Paragraph(f"<b>User:</b> {safe_user_name}", body_style),
        Paragraph(f"<b>Range:</b> {range_label}", body_style),
        Paragraph(f"<b>Total spent:</b> ₹{total:,.2f}", body_style),
        Spacer(1, 0.5 * cm),
    ]

    if not expenses:
        story.append(Paragraph("No expenses in this date range.", empty_style))
    else:
        header = ["Date", "Description", "Category", "Amount"]
        rows = [header]
        for e in expenses:
            rows.append(
                [
                    e["date"],
                    e["description"],
                    e["category"],
                    f"₹{e['amount']:,.2f}",
                ]
            )
        table = Table(
            rows,
            colWidths=[2.6 * cm, 7 * cm, 3.2 * cm, 3.2 * cm],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), PDF_FONT_BOLD),
                    ("FONTNAME", (0, 1), (-1, -1), PDF_FONT_REGULAR),
                    ("FONTSIZE", (0, 0), (-1, -1), 10),
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(PDF_COLOR_HEADER_BG),
                    ),
                    ("ALIGN", (3, 0), (3, -1), "RIGHT"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor(PDF_COLOR_ROW_ALT)],
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.25,
                        colors.HexColor(PDF_COLOR_HEADER_BG),
                    ),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)

    story.append(Spacer(1, 0.6 * cm))
    story.append(
        Paragraph(
            f"Generated {datetime.now():%Y-%m-%d %H:%M}",
            ParagraphStyle(
                "Footer",
                parent=body_style,
                fontSize=9,
                textColor=colors.HexColor(PDF_COLOR_MUTED),
            ),
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer


def _build_xlsx(expenses):
    """Render the expense list as an .xlsx workbook and return a BytesIO."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"

    headers = ["Date", "Description", "Category", "Amount (INR)"]
    ws.append(headers)
    header_font = Font(bold=True)
    header_fill = PatternFill(
        start_color=XLSX_COLOR_HEADER_BG,
        end_color=XLSX_COLOR_HEADER_BG,
        fill_type="solid",
    )
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")
    ws.cell(row=1, column=4).alignment = Alignment(horizontal="right")

    for e in expenses:
        ws.append(
            [
                e["date"],
                _safe_xlsx_string(e["description"]),
                e["category"],
                float(e["amount"]),
            ]
        )

    last_data_row = len(expenses) + 1  # +1 for header row
    if expenses:
        for row in range(2, last_data_row + 1):
            ws.cell(row=row, column=4).number_format = "#,##0.00"
            ws.cell(row=row, column=4).alignment = Alignment(horizontal="right")

        total_row = last_data_row + 1
        ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
        formula_cell = ws.cell(
            row=total_row,
            column=4,
            value=f"=SUM(D2:D{last_data_row})",
        )
        formula_cell.font = Font(bold=True)
        formula_cell.number_format = "#,##0.00"
        formula_cell.alignment = Alignment(horizontal="right")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ------------------------------------------------------------------ #
# Routes                                                              #
# ------------------------------------------------------------------ #


@app.route("/")
def landing():
    if session.get("user_id"):
        return redirect(url_for("profile"))
    return render_template("landing.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if session.get("user_id"):
        return redirect(url_for("landing"))

    if request.method == "GET":
        return render_template("register.html")

    name = request.form.get("name", "").strip()
    email = request.form.get("email", "").strip()
    password = request.form.get("password", "")
    confirm_password = request.form.get("confirm_password", "")

    if not name or not email or not password or not confirm_password:
        return render_template(
            "register.html", error="All fields are required.", name=name, email=email
        )

    if len(password) < MIN_PASSWORD_LENGTH:
        return render_template(
            "register.html",
            error=f"Password must be at least {MIN_PASSWORD_LENGTH} characters.",
            name=name,
            email=email,
        )

    if password != confirm_password:
        return render_template(
            "register.html", error="Passwords do not match.", name=name, email=email
        )

    try:
        create_user(name, email, password)
    except sqlite3.IntegrityError:
        return render_template(
            "register.html", error="Email already registered.", name=name, email=email
        )

    flash("Account created successfully! Please sign in.")
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return redirect(url_for("landing"))

    if request.method == "GET":
        return render_template("login.html")

    email = request.form.get("email", "").strip()
    password = request.form.get("password", "")

    if not email or not password:
        return render_template(
            "login.html", error="All fields are required.", email=email
        )

    user = get_user_by_email(email)
    if user is None or not check_password_hash(user["password_hash"], password):
        return render_template(
            "login.html", error="Invalid email or password.", email=email
        )

    session.clear()
    session["user_id"] = user["id"]
    session["user_name"] = user["name"]
    return redirect(url_for("landing"))


# ------------------------------------------------------------------ #
# Authenticated routes                                                #
# ------------------------------------------------------------------ #


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("landing"))


@app.route("/profile")
@login_required
def profile():
    date_from, date_to, range_error = _resolve_date_range_args(request)
    if range_error:
        flash(range_error, "error")

    presets = _resolve_presets(date.today())
    if date_from and date_to:
        active_preset = next(
            (key for key, bounds in presets.items() if bounds == (date_from, date_to)),
            "custom",
        )
    else:
        active_preset = "all"

    preset_labels = {
        "all": "all time",
        "this_month": "this month",
        "last_3_months": "last 3 months",
        "last_6_months": "last 6 months",
    }
    filter_label = preset_labels.get(active_preset) or f"{date_from} – {date_to}"

    user = get_user_by_id(session["user_id"])
    if user is None:
        session.clear()
        return redirect(url_for("login"))

    stats = get_summary_stats(session["user_id"], date_from=date_from, date_to=date_to)
    transactions = get_recent_transactions(
        session["user_id"], date_from=date_from, date_to=date_to
    )
    categories = get_category_breakdown(
        session["user_id"], date_from=date_from, date_to=date_to
    )

    return render_template(
        "profile.html",
        user=user,
        stats=stats,
        transactions=transactions,
        categories=categories,
        presets=presets,
        active_preset=active_preset,
        filter_label=filter_label,
        date_from=date_from or "",
        date_to=date_to or "",
    )


@app.route("/analytics")
@login_required
def analytics():
    return render_template("analytics.html")


@app.route("/expenses/add", methods=["GET", "POST"])
@login_required
def add_expense():
    if request.method == "GET":
        return render_template(
            "add_expense.html",
            categories=CATEGORIES,
            today=date.today().isoformat(),
        )

    fields = _expense_form_fields()
    amount, error = _validate_expense_form(
        fields["amount"], fields["category"], fields["date"]
    )
    if error:
        return render_template(
            "add_expense.html",
            categories=CATEGORIES,
            today=date.today().isoformat(),
            error=error,
            **fields,
        )

    insert_expense(
        session["user_id"],
        amount,
        fields["category"],
        fields["date"],
        fields["description"],
    )
    flash("Expense added.")
    return redirect(url_for("profile"))


@app.route("/expenses/<int:expense_id>/edit", methods=["GET", "POST"])
@login_required
def edit_expense(expense_id):
    expense = get_expense_by_id(expense_id, session["user_id"])
    if expense is None:
        flash("Expense not found.", "error")
        return redirect(url_for("profile"))

    if request.method == "GET":
        return render_template(
            "edit_expense.html",
            expense=expense,
            categories=CATEGORIES,
            amount=expense["amount"],
            category=expense["category"],
            date=expense["date"],
            description=expense["description"] or "",
        )

    fields = _expense_form_fields()
    amount, error = _validate_expense_form(
        fields["amount"], fields["category"], fields["date"]
    )
    if error:
        return render_template(
            "edit_expense.html",
            expense=expense,
            categories=CATEGORIES,
            error=error,
            **fields,
        )

    rows_affected = update_expense(
        expense_id,
        session["user_id"],
        amount,
        fields["category"],
        fields["date"],
        fields["description"],
    )
    if rows_affected == 0:
        flash("Expense not found.", "error")
        return redirect(url_for("profile"))
    flash("Expense updated.")
    return redirect(url_for("profile"))


@app.route("/expenses/<int:expense_id>/delete", methods=["POST"])
@login_required
def delete_expense(expense_id):
    rows_affected = delete_expense_query(expense_id, session["user_id"])
    if rows_affected == 0:
        flash("Expense not found.", "error")
    else:
        flash("Expense deleted.")
    return redirect(url_for("profile"))


@app.route("/expenses/export/pdf")
@login_required
def export_expenses_pdf():
    user = get_user_by_id(session["user_id"])
    if user is None:
        session.clear()
        return redirect(url_for("login"))

    date_from, date_to, _ = _resolve_date_range_args(request)
    expenses = get_expenses_for_export(session["user_id"], date_from, date_to)

    buffer = _build_pdf(user["name"], _human_range_label(date_from, date_to), expenses)
    filename = (
        f"spendly-{_slugify_name(user['name'])}-"
        f"{_filename_range_label(date_from, date_to)}.pdf"
    )
    return send_file(
        buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=filename,
    )


@app.route("/expenses/export/xlsx")
@login_required
def export_expenses_xlsx():
    user = get_user_by_id(session["user_id"])
    if user is None:
        session.clear()
        return redirect(url_for("login"))

    date_from, date_to, _ = _resolve_date_range_args(request)
    expenses = get_expenses_for_export(session["user_id"], date_from, date_to)

    buffer = _build_xlsx(expenses)
    filename = (
        f"spendly-{_slugify_name(user['name'])}-"
        f"{_filename_range_label(date_from, date_to)}.xlsx"
    )
    return send_file(
        buffer,
        mimetype=XLSX_MIMETYPE,
        as_attachment=True,
        download_name=filename,
    )


@app.route("/terms")
def terms():
    return render_template("terms.html")


@app.route("/privacy")
def privacy():
    return render_template("privacy.html")


if __name__ == "__main__":
    debug = os.environ.get("FLASK_DEBUG", "false").lower() in ("1", "true", "yes")
    app.run(debug=debug, port=5001)
